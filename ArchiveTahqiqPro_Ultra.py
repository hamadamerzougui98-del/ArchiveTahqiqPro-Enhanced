# -*- coding: utf-8 -*-
import sys
import os
import re
import html
import time
import shutil
import sqlite3
import json
import hashlib
import secrets
import zipfile
import tempfile
import traceback
from datetime import date

from PyQt5.QtCore import (
    Qt, QDate, QTimer, QMarginsF, QSettings,
    QPropertyAnimation, QEasingCurve, QSortFilterProxyModel
)
from PyQt5.QtGui import QColor, QFont, QFontDatabase, QTextDocument, QTextTable, QPageLayout, QPageSize, QKeySequence
from PyQt5.QtWidgets import *
from PyQt5.QtPrintSupport import QPrinter

ORG = 'DCCF Djelfa'
DATA_DIR = os.path.join(os.environ.get('APPDATA', '.'), 'ArchiveTahqiq')
ATT_DIR = os.path.join(DATA_DIR, 'attachments')
AUTH = os.path.join(DATA_DIR, 'auth.json')

PROV = ['الجلفة', 'عين وسارة', 'حاسي بحبح', 'الإدريسية', 'مسعد']
HEAD = ['العدد', 'رقم الملف', 'الجهة المرسلة', 'رقم الإرسال', 'تاريخ الإرسال', *PROV, 'الحالة']
COLS = HEAD + ['المرفق']
C_PROV, C_STATUS, C_ATT = 5, 10, 11
YES, NO = '✔', '✘'
MAX_HIST = 100

SENDERS = [
    'بلدية الجلفة', 'بلدية عين وسارة', 'بلدية حاسي بحبح', 'بلدية الإدريسية', 'بلدية مسعد',
    'مديرية مسح الأراضي', 'مديرية أملاك الدولة', 'المحكمة', 'الولاية'
]

GREEN_BG, RED_BG = QColor('#dff0e5'), QColor('#fbe9e8')
GREEN_FG, RED_FG = QColor('#136b40'), QColor('#b52a1c')
DONE_BG, DONE_FG = QColor('#d6ecdf'), QColor('#0f4a33')
PEND_BG, PEND_FG = QColor('#fbeedb'), QColor('#6e4410')
HL = QColor('#ffe9a8')

def load_arabic_font():
    candidates = [
        'Tahoma', 'Segoe UI', 'Arial', 'Noto Sans Arabic', 'Cairo', 'Janna', 'Aref Ruqaa', 'Amiri', 'Scheherazade',
        'Arial Unicode MS', 'Microsoft Sans Serif'
    ]
    db = QFontDatabase()
    for name in candidates:
        if db.families().__contains__(name):
            return QFont(name, 10)
    return QFont()

APP_FONT = load_arabic_font()

QSS = """
*{font-family:'Segoe UI',Tahoma,'Noto Sans Arabic','Cairo';font-size:15px;font-weight:600}
QWidget{background:#f4efe8;color:#171717}
QMainWindow,QDialog,QWidget{background:#f4efe8}
#hdr{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #0b221d,stop:0.5 #123c2f,stop:1 #3a4d2d);border-radius:18px}
#hdr QLabel{background:transparent;color:white}
#hdr QLabel#title{font-size:26px;font-weight:700}
#hdr QLabel#gold{color:#e7cf9c;font-size:13px;font-weight:600}
QPushButton#hdrBtn{background:rgba(255,255,255,0.10);color:white;border:1px solid #d4ae73;border-radius:10px;padding:6px 12px}
QPushButton#hdrBtn:hover{background:rgba(255,255,255,0.18)}
QFrame#card{background:linear-gradient(180deg,#ffffff 0%,#f8f2e6 100%);border:1px solid #e0d3b6;border-radius:14px}
QLabel#big{font-size:22px;font-weight:700;color:#0a2a21}
QPushButton{background:#fff;border:1px solid #d9c9a8;border-radius:10px;padding:7px 14px;color:#111}
QPushButton:hover{border-color:#b58843;background:#fffaf4}
QPushButton#addBtn{background:qlineargradient(x1:0,y1:0,x2:0,y2:1,stop:0 #d7aa66,stop:1 #a97a3f);color:white;border:2px solid #f1dfba;border-radius:16px;font-size:21px;font-weight:bold;padding:9px 26px}
QPushButton#addBtn:hover{background:#b98741}
QPushButton#tab{border-radius:17px;padding:7px 18px}
QPushButton#tab:checked{background:#113d32;color:white;border-color:#113d32}
QPushButton#hist{font-size:24px;font-weight:bold;color:#123b2e;padding:2px 14px;min-width:44px;border:2px solid #123b2e;border-radius:12px}
QPushButton#hist:hover{background:#e4f2ea}
QPushButton#hist:disabled{color:#bdb3a0;border-color:#ddd2bd;background:#f1ebe0}
QPushButton#fab{background:#143f34;color:white;border:3px solid #d9b67a;border-radius:22px;padding:9px 42px;font-size:17px;font-weight:bold}
QPushButton#fab:hover{background:#205746}
QLineEdit,QComboBox,QDateEdit{background:white;border:1px solid #d8c8a8;border-radius:10px;padding:6px 10px;min-height:22px;color:#111}
QLineEdit:focus,QComboBox:focus,QDateEdit:focus{border:2px solid #b8894a}
QLineEdit#search{font-size:17px;padding:8px 14px;border-radius:12px}
QProgressBar{border:1px solid #d7cab5;border-radius:8px;background:#f7f3ee;text-align:center;min-height:22px;color:#111;font-weight:bold}
QProgressBar::chunk{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #d1a15b,stop:1 #4f9a78);border-radius:7px}
QTableWidget{background:white;alternate-background-color:#fbf8f2;border:1px solid #ddd2bd;gridline-color:#ebdfc8;color:#111}
QHeaderView::section{background:#123b2e;color:white;padding:9px;border:none;border-left:1px solid #1a5140;font-weight:bold;font-size:15px}
QTableWidget::item:selected{background:#f3e3c7;color:black}
QGroupBox{background:white;border:1px solid #ddd2bd;border-radius:10px;margin-top:16px;padding:12px}
QGroupBox::title{subcontrol-origin:margin;right:12px;color:#123b2e;font-weight:bold}
QCheckBox{background:transparent;font-size:16px;spacing:8px}
QCheckBox::indicator{width:24px;height:24px}
QLabel#hint{color:#2f2a22}
QLabel#msg{color:#0f4a33;background:#e4f2ea;border-radius:8px;padding:5px 12px}
QMenu{background:white;border:1px solid #d6c9b1;padding:6px}
QMenu::item{padding:8px 26px}
QMenu::item:selected{background:#f3e3c7;color:black}
QDialog{background:#f4efe8}
QAbstractScrollArea{border:0px solid transparent}
"""

DARK_QSS = """
*{font-family:'Segoe UI',Tahoma,'Noto Sans Arabic','Cairo';font-size:15px;font-weight:600}
QWidget{background:#101615;color:#edf5f2}
QMainWindow,QDialog,QWidget{background:#101615}
#hdr{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #081610,stop:0.5 #123126,stop:1 #2d3b2d);border-radius:18px}
#hdr QLabel{background:transparent;color:#f5f7f5}
#hdr QLabel#title{font-size:26px;font-weight:700}
#hdr QLabel#gold{color:#e7cf9c;font-size:13px;font-weight:600}
QPushButton#hdrBtn{background:rgba(255,255,255,0.08);color:white;border:1px solid #cba66c;border-radius:10px;padding:6px 12px}
QPushButton#hdrBtn:hover{background:rgba(255,255,255,0.16)}
QFrame#card{background:linear-gradient(180deg,#1f2826 0%,#171e1d 100%);border:1px solid #3a4742;border-radius:14px}
QLabel#big{font-size:22px;font-weight:700;color:#dfeee7}
QPushButton{background:#1d2422;border:1px solid #48554f;border-radius:10px;padding:7px 14px;color:#edf3ef}
QPushButton:hover{border-color:#c29a60;background:#2a332f}
QPushButton#addBtn{background:qlineargradient(x1:0,y1:0,x2:0,y2:1,stop:0 #c89a58,stop:1 #8d642f);color:white;border:2px solid #e3c48f;border-radius:16px;font-size:21px;font-weight:bold;padding:9px 26px}
QPushButton#addBtn:hover{background:#b17d3f}
QPushButton#tab{border-radius:17px;padding:7px 18px}
QPushButton#tab:checked{background:#c39a60;color:#101513;border-color:#c39a60}
QPushButton#hist{font-size:24px;font-weight:bold;color:#d5e8dc;padding:2px 14px;min-width:44px;border:2px solid #4b6b5b;border-radius:12px}
QPushButton#hist:hover{background:#293a33}
QPushButton#hist:disabled{color:#6e7873;border-color:#39423e;background:#1a201e}
QPushButton#fab{background:#d0a365;color:#111814;border:3px solid #ead09e;border-radius:22px;padding:9px 42px;font-size:17px;font-weight:bold}
QPushButton#fab:hover{background:#e0b879}
QLineEdit,QComboBox,QDateEdit{background:#1d2421;border:1px solid #46524d;border-radius:10px;padding:6px 10px;min-height:22px;color:#f0f5f2}
QLineEdit:focus,QComboBox:focus,QDateEdit:focus{border:2px solid #c49a61}
QLineEdit#search{font-size:17px;padding:8px 14px;border-radius:12px}
QProgressBar{border:1px solid #46524d;border-radius:8px;background:#1d2421;text-align:center;min-height:22px;color:#edf3ef;font-weight:bold}
QProgressBar::chunk{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #b99055,stop:1 #4d9875);border-radius:7px}
QTableWidget{background:#1b211f;alternate-background-color:#222a27;border:1px solid #3b4742;gridline-color:#303a36;color:#edf3ef}
QHeaderView::section{background:#244c3d;color:white;padding:9px;border:none;border-left:1px solid #32604f;font-weight:bold;font-size:15px}
QTableWidget::item:selected{background:#66522f;color:white}
QGroupBox{background:#202725;border:1px solid #3d4944;border-radius:10px;margin-top:16px;padding:12px}
QGroupBox::title{subcontrol-origin:margin;right:12px;color:#d9bb83;font-weight:bold}
QCheckBox{background:transparent;font-size:16px;spacing:8px}
QCheckBox::indicator{width:24px;height:24px}
QLabel#hint{color:#aebbb5}
QLabel#msg{color:#d9eee2;background:#234b3b;border-radius:8px;padding:5px 12px}
QMenu{background:#202725;color:#edf3ef;border:1px solid #46524d;padding:6px}
QMenu::item{padding:8px 26px}
QMenu::item:selected{background:#3a4b43;color:white}
QDialog{background:#101615}
QAbstractScrollArea{border:0px solid transparent}
QScrollBar:vertical,QScrollBar:horizontal{background:#171c1a;border:none;margin:2px}
QScrollBar::handle:vertical,QScrollBar::handle:horizontal{background:#53635c;border-radius:6px;min-height:30px}
QToolTip{background:#111615;color:#f2f6f4;border:1px solid #c49a61;padding:6px}
"""


def _excepthook(t, v, tb):
    msg = ''.join(traceback.format_exception(t, v, tb))
    try:
        os.makedirs(DATA_DIR, exist_ok=True)
        with open(os.path.join(DATA_DIR, 'errors.log'), 'a', encoding='utf-8') as f:
            f.write(f'\n[{time.strftime("%Y-%m-%d %H:%M:%S")}\n{msg}')
    except Exception:
        pass
    QMessageBox.critical(None, 'خطأ', f'حدث خطأ (لم يُغلق البرنامج):\n\n{v}\n\nالتفاصيل محفوظة في ملف errors.log')

sys.excepthook = _excepthook


def done(r):
    return sum(1 for i in range(5) if r.get(f'p{i}'))


def iso2d(s):
    return f'{s[8:10]}/{s[5:7]}/{s[:4]}' if s and len(s) >= 10 else ''


def to_iso(v):
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    v = str(v or '').strip()
    m = re.match(r'(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})', v)
    if m:
        return f'{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}'
    m = re.match(r'(\d{4})-(\d{1,2})-(\d{1,2})', v)
    return f'{m[1]}-{int(m[2]):02d}-{int(m[3]):02d}' if m else ''


def norm(s):
    s = str(s or '').lower()
    s = re.sub('[\u064B-\u0652\u0640]', '', s)
    s = re.sub('[أإآٱ]', 'ا', s)
    return s.replace('ة', 'ه').replace('ى', 'ي').replace('ؤ', 'و').replace('ئ', 'ي')


def nat(s):
    return [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', norm(s))]


def serial(s):
    nums = [int(x) for x in re.findall(r'\d+', str(s))]
    if not nums:
        return None
    other = [n for n in nums if not 1900 <= n <= 2100]
    return other[0] if other else nums[-1]


def in_range(fno, lo, hi):
    bounds = [b for b in (lo, hi) if b]
    if all(b.isdigit() for b in bounds):
        v = serial(fno)
        return v is not None and (not lo or v >= int(lo)) and (not hi or v <= int(hi))
    try:
        k = nat(fno)
        return (not lo or k >= nat(lo)) and (not hi or k <= nat(hi))
    except TypeError:
        return False


def truthy(v):
    v = str(v if v is not None else '').strip()
    return bool(v) and v not in (NO, 'x', 'X', 'لا', '0', 'False', 'None') and 'بانتظار' not in v


def dateEdit(iso=''):
    d = QDateEdit(calendarPopup=True)
    d.setDisplayFormat('dd/MM/yyyy')
    d.setDate(QDate.fromString(iso, 'yyyy-MM-dd') if iso else QDate.currentDate())
    return d


def combo(items):
    c = QComboBox()
    c.addItems(items)
    return c


PDF_HEAD = ['العدد', 'رقم الملف', 'الجهة المرسلة', 'رقم الإرسال', *PROV, 'الحالة']


def table_doc(L, title, selected_prov=None):
    heads = ['العدد', 'رقم الملف', 'الجهة المرسلة', 'رقم الإرسال']
    prov_indexes = list(range(5)) if selected_prov is None else [selected_prov]
    heads += [PROV[k] for k in prov_indexes]
    heads += ['الحالة']
    e = html.escape
    dn = sum(done(r) == 5 for r in L)

    th = ''.join(
        f'<th align="center" bgcolor="#123b2e" style="color:white;font-weight:bold;padding:3px;font-size:8pt">{e(x)}</th>'
        for x in heads
    )

    trs = []
    for i, r in enumerate(L):
        row = [str(i + 1), r.get('file_no', ''), r.get('sender', ''), r.get('send_no', '')]
        tds = ''.join(
            f'<td dir="rtl" align="right" style="padding:3px;font-weight:bold;font-size:8pt">{e(str(v))}</td>'
            for v in row
        )
        for k in prov_indexes:
            if r.get(f'p{k}'):
                tds += '<td dir="rtl" align="center" bgcolor="#dff0e5" style="padding:2px;color:#136b40;font-size:11pt;font-weight:bold">✔</td>'
            else:
                tds += '<td dir="rtl" align="center" bgcolor="#fbe9e8" style="padding:2px;color:#b52a1c;font-size:11pt;font-weight:bold">✘</td>'

        n = done(r)
        bg = '#d6ecdf' if n == 5 else '#fbeedb'
        status = 'مكمل الإجابة' if n == 5 else f'معلق {n}/5'
        tds += f'<td dir="rtl" align="center" bgcolor="{bg}" style="padding:3px;font-weight:bold;font-size:7.5pt">{e(status)}</td>'
        trs.append(f'<tr>{tds}</tr>')

    html_content = (
        '<html dir="rtl"><head><meta charset="utf-8"></head>'
        '<body dir="rtl" style="font-family:Tahoma;font-size:8pt;color:#000">'
        f'<h2 align="center" style="margin:0 0 5px;color:#123b2e">{e(title)}</h2>'
        f'<p align="center" style="margin:0 0 7px;font-size:8pt"><b>'
        f'عدد الملفات: {len(L)} &nbsp; | &nbsp; '
        f'مكمل الإجابة: {dn} &nbsp; | &nbsp; '
        f'معلق: {len(L)-dn} &nbsp; | &nbsp; '
        f'التاريخ: {date.today():%d/%m/%Y}'
        '</b></p>'
        '<table dir="rtl" border="1" cellspacing="0" cellpadding="0" width="100%" style="border-collapse:collapse;table-layout:fixed">'
        f'<thead><tr>{th}</tr></thead>'
        f'{"".join(trs)}'
        '</table>'
        f'<p align="left" style="font-size:7pt;margin-top:5px"><small>{e(ORG)}</small></p>'
        '</body></html>'
    )

    doc = QTextDocument()
    doc.setHtml(html_content)
    for fr in doc.rootFrame().childFrames():
        if isinstance(fr, QTextTable):
            tf = fr.format()
            tf.setLayoutDirection(Qt.RightToLeft)
            tf.setHeaderRowCount(1)
            fr.setFormat(tf)
    return doc


def a4_printer(landscape=False):
    pr = QPrinter(QPrinter.HighResolution)
    pr.setPageLayout(
        QPageLayout(
            QPageSize(QPageSize.A4),
            QPageLayout.Landscape if landscape else QPageLayout.Portrait,
            QMarginsF(7, 7, 7, 7),
            QPageLayout.Millimeter,
        )
    )
    return pr


class SItem(QTableWidgetItem):
    def __init__(self, text, key=None):
        super().__init__(str(text))
        self.key = key if key is not None else str(text)
        self.setTextAlignment(Qt.AlignCenter)
        self.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)

    def __lt__(self, o):
        try:
            return self.key < o.key
        except TypeError:
            return str(self.key) < str(o.key)


class DB:
    def __init__(self):
        os.makedirs(ATT_DIR, exist_ok=True)
        self.c = sqlite3.connect(os.path.join(DATA_DIR, 'archive.db'))
        self.c.row_factory = sqlite3.Row
        cols = ','.join(f"p{i} TEXT DEFAULT '', t{i} TEXT DEFAULT ''" for i in range(5))
        self.c.execute(
            f"CREATE TABLE IF NOT EXISTS files("
            f"id INTEGER PRIMARY KEY AUTOINCREMENT,"
            f"file_no TEXT UNIQUE,"
            f"sender TEXT DEFAULT '',"
            f"send_no TEXT DEFAULT '',"
            f"send_date TEXT DEFAULT '',"
            f"{cols},"
            f"att TEXT DEFAULT '',"
            f"notes TEXT DEFAULT ''"
            f")"
        )
        self.c.commit()

    @staticmethod
    def _d(r):
        return {k: (v if v is not None else '') for k, v in dict(r).items()}

    def all(self):
        return [self._d(r) for r in self.c.execute('SELECT * FROM files ORDER BY id')]

    def get(self, i):
        r = self.c.execute('SELECT * FROM files WHERE id=?', (i,)).fetchone()
        return self._d(r) if r else None

    def by_no(self, no):
        r = self.c.execute('SELECT id FROM files WHERE file_no=?', (no,)).fetchone()
        return r[0] if r else None

    def save(self, d, i=None, commit=True):
        if i:
            self.c.execute(
                f'UPDATE files SET {",".join(k + "=?" for k in d)} WHERE id=?',
                [*d.values(), i]
            )
        else:
            i = self.c.execute(
                f'INSERT INTO files({",".join(d)}) VALUES({",".join("?" * len(d))})',
                list(d.values())
            ).lastrowid
        if commit:
            self.c.commit()
        return i

    def delete(self, i):
        self.c.execute('DELETE FROM files WHERE id=?', (i,))
        self.c.commit()

    def raw_delete(self, i):
        self.c.execute('DELETE FROM files WHERE id=?', (i,))

    def put(self, row):
        self.c.execute(
            f'INSERT OR REPLACE INTO files({",".join(row)}) VALUES({",".join("?" * len(row))})',
            list(row.values())
        )


def _hash(pw, salt):
    return hashlib.pbkdf2_hmac('sha256', pw.encode('utf-8'), bytes.fromhex(salt), 200000).hex()


def has_pw():
    return os.path.exists(AUTH)


def set_pw(pw):
    os.makedirs(DATA_DIR, exist_ok=True)
    salt = secrets.token_hex(16)
    with open(AUTH, 'w', encoding='utf-8') as f:
        json.dump({'salt': salt, 'hash': _hash(pw, salt)}, f)


def check_pw(pw):
    try:
        with open(AUTH, encoding='utf-8') as f:
            d = json.load(f)
        return secrets.compare_digest(_hash(pw, d['salt']), d['hash'])
    except Exception:
        return False


def pwedit(ph):
    e = QLineEdit(placeholderText=ph)
    e.setEchoMode(QLineEdit.Password)
    e.setMinimumHeight(40)
    return e


def show_toggle(fields):
    c = QCheckBox('إظهار كلمة المرور')
    c.toggled.connect(lambda v: [f.setEchoMode(QLineEdit.Normal if v else QLineEdit.Password) for f in fields])
    return c


class LoginDlg(QDialog):
    def __init__(self, first=False):
        super().__init__(None, Qt.WindowCloseButtonHint)
        self.first = first
        self.tries = 0
        self.setWindowTitle('أرشيف التحقيق العقاري — تسجيل الدخول')
        self.setFixedWidth(470)
        self.setStyleSheet(QSS if not self.is_dark() else DARK_QSS)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 18)
        lay.setSpacing(10)

        hdr = QFrame(objectName='hdr')
        hl = QVBoxLayout(hdr)
        hl.setContentsMargins(20, 16, 20, 16)
        for t, n in [('🔒', 'title'), ('أرشيف التحقيق العقاري', 'title'), (f'ولاية الجلفة  —  {ORG}', 'gold')]:
            l = QLabel(t, objectName=n)
            l.setAlignment(Qt.AlignCenter)
            hl.addWidget(l)
        lay.addWidget(hdr)

        body = QVBoxLayout()
        body.setContentsMargins(26, 6, 26, 0)
        body.setSpacing(10)
        lay.addLayout(body)

        info = QLabel('مرحباً 👋  أنشئ كلمة مرور لحماية البرنامج (4 أحرف على الأقل)' if first else 'أدخل كلمة المرور للدخول')
        info.setAlignment(Qt.AlignCenter)
        info.setWordWrap(True)
        body.addWidget(info)

        self.pw = pwedit('كلمة المرور')
        body.addWidget(self.pw)
        self.pw2 = pwedit('تأكيد كلمة المرور')
        if first:
            body.addWidget(self.pw2)
            self.pw.returnPressed.connect(self.pw2.setFocus)
        body.addWidget(show_toggle([self.pw, self.pw2]))

        self.err = QLabel()
        self.err.setStyleSheet('color:#b52a1c;font-weight:bold')
        self.err.setAlignment(Qt.AlignCenter)
        body.addWidget(self.err)

        self.btn = QPushButton('إنشاء كلمة المرور والدخول' if first else '🔓  دخول', objectName='addBtn')
        self.btn.setDefault(True)
        self.btn.setCursor(Qt.PointingHandCursor)
        self.btn.clicked.connect(self.go)
        body.addWidget(self.btn)

        self.pw.setFocus()

    @staticmethod
    def is_dark():
        return False

    def fail(self, m):
        self.err.setText(m)
        self.pw.setFocus()

    def go(self):
        p = self.pw.text()
        if self.first:
            if len(p) < 4:
                return self.fail('كلمة المرور قصيرة (4 أحرف على الأقل)')
            if p != self.pw2.text():
                return self.fail('كلمتا المرور غير متطابقتين')
            set_pw(p)
            return self.accept()
        if check_pw(p):
            return self.accept()
        self.tries += 1
        self.pw.clear()
        if self.tries >= 5:
            self.tries = 0
            self.fail('⛔ محاولات خاطئة كثيرة — انتظر 30 ثانية')
            self.pw.setEnabled(False)
            self.btn.setEnabled(False)
            QTimer.singleShot(30000, lambda: (self.pw.setEnabled(True), self.btn.setEnabled(True), self.err.clear(), self.pw.setFocus()))
        else:
            self.fail(f'كلمة المرور خاطئة ✘   (المحاولة {self.tries} من 5)')


class ChangePwDlg(QDialog):
    def __init__(self, parent):
        super().__init__(parent)
        self.setWindowTitle('🔑 تغيير كلمة المرور')
        self.setFixedWidth(480)
        lay = QVBoxLayout(self)
        f = QFormLayout()
        f.setVerticalSpacing(10)
        self.old = pwedit('الحالية')
        self.new = pwedit('الجديدة (4 أحرف على الأقل)')
        self.new2 = pwedit('أعد كتابة الجديدة')
        f.addRow('كلمة المرور الحالية', self.old)
        f.addRow('كلمة المرور الجديدة', self.new)
        f.addRow('تأكيد الجديدة', self.new2)
        lay.addLayout(f)
        lay.addWidget(show_toggle([self.old, self.new, self.new2]))
        self.err = QLabel()
        self.err.setStyleSheet('color:#b52a1c;font-weight:bold')
        lay.addWidget(self.err)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Save).setText('💾 حفظ')
        bb.button(QDialogButtonBox.Cancel).setText('إلغاء')
        bb.accepted.connect(self.ok)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)

    def ok(self):
        if not check_pw(self.old.text()):
            self.err.setText('كلمة المرور الحالية خاطئة ✘')
            return self.old.setFocus()
        if len(self.new.text()) < 4:
            self.err.setText('كلمة المرور الجديدة قصيرة (4 أحرف على الأقل)')
            return self.new.setFocus()
        if self.new.text() != self.new2.text():
            self.err.setText('كلمتا المرور الجديدتان غير متطابقتين')
            return self.new2.setFocus()
        set_pw(self.new.text())
        QMessageBox.information(self, 'تم', 'تم تغيير كلمة المرور بنجاح ✓')
        self.accept()


class FileDlg(QDialog):
    def __init__(self, parent, senders, r=None):
        super().__init__(parent)
        self.r = r or {}
        self.setWindowTitle(f'تعديل الملف {r["file_no"]}' if r else 'إضافة ملف جديد')
        self.setMinimumWidth(760)
        self.att = r.get('att', '')
        self.no = QLineEdit(r.get('file_no', ''), placeholderText='2025/096')
        self.sender = QComboBox(editable=True)
        self.sender.addItems(senders)
        self.sender.setCurrentText(r.get('sender', ''))
        self.sno = QLineEdit(r.get('send_no', ''), placeholderText='55/2025')
        self.sdate = dateEdit(r.get('send_date', ''))
        self.notes = QLineEdit(r.get('notes', ''))

        lay = QVBoxLayout(self)
        f = QFormLayout()
        f.setVerticalSpacing(10)
        for t, w in [('رقم الملف *', self.no), ('الجهة المرسلة *', self.sender), ('رقم الإرسال', self.sno), ('تاريخ الإرسال', self.sdate)]:
            f.addRow(t, w)
        lay.addLayout(f)

        g = QGroupBox('ردود المحافظات العقارية   (✔ أجابت  ·  ✘ لم تُجب)')
        gl = QHBoxLayout(g)
        self.chk = []
        for i, n in enumerate(PROV):
            c = QCheckBox(n)
            c.setChecked(bool(r.get(f'p{i}')))
            gl.addWidget(c)
            self.chk.append(c)
        ab = QPushButton('تحديد / إلغاء الكل')
        ab.clicked.connect(self.toggle_all)
        gl.addWidget(ab)
        lay.addWidget(g)

        self.attBtn = QPushButton()
        self.attBtn.setMinimumHeight(44)
        self.attBtn.clicked.connect(self.pick)
        self.upd()
        lay.addWidget(self.attBtn)

        f2 = QFormLayout()
        f2.addRow('ملاحظات', self.notes)
        lay.addLayout(f2)

        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Save).setText('💾 حفظ')
        bb.button(QDialogButtonBox.Cancel).setText('إلغاء')
        bb.accepted.connect(self.ok)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)

    def toggle_all(self):
        v = not all(c.isChecked() for c in self.chk)
        for c in self.chk:
            c.setChecked(v)

    def upd(self):
        if self.att:
            self.attBtn.setText(f'📎 {self.att.split("_", 1)[-1]}   (تغيير)')
        else:
            self.attBtn.setText('📎 إرفاق ملف الجهة المرسلة (PDF / صورة / Word)')

    def pick(self):
        p, _ = QFileDialog.getOpenFileName(self, 'مرفق الجهة المرسلة', '', 'مستندات (*.pdf *.jpg *.jpeg *.png *.doc *.docx *.tif);;كل الملفات (*)')
        if p:
            self.att = f'{int(time.time())}_{os.path.basename(p)}'
            shutil.copy2(p, os.path.join(ATT_DIR, self.att))
            self.upd()

    def ok(self):
        if not self.no.text().strip() or not self.sender.currentText().strip():
            return QMessageBox.warning(self, 'تنبيه', 'رقم الملف والجهة المرسلة إجباريان')
        self.accept()

    def data(self):
        d = {
            'file_no': self.no.text().strip(),
            'sender': self.sender.currentText().strip(),
            'send_no': self.sno.text().strip(),
            'send_date': self.sdate.date().toString('yyyy-MM-dd'),
            'att': self.att,
            'notes': self.notes.text().strip(),
        }
        today = date.today().isoformat()
        for i, c in enumerate(self.chk):
            on = c.isChecked()
            d[f'p{i}'] = (self.r.get(f'p{i}') or '1') if on else ''
            d[f't{i}'] = (self.r.get(f't{i}') or today) if on else ''
        return d


class BulkActionDlg(QDialog):
    def __init__(self, parent, count):
        super().__init__(parent)
        self.setWindowTitle('⚡ إجراء جماعي')
        self.setFixedWidth(450)
        lay = QVBoxLayout(self)
        lay.setSpacing(12)
        info = QLabel(f'سيتم تطبيق الإجراء على <b style="color:#123b2e;font-size:18px">{count}</b> ملف/ملفات ظاهرة حالياً في الجدول.')
        info.setAlignment(Qt.AlignCenter)
        info.setWordWrap(True)
        lay.addWidget(info)

        f = QFormLayout()
        f.setVerticalSpacing(10)
        self.prov = QComboBox()
        self.prov.addItems(['كل المحافظات'] + PROV)
        f.addRow('المحافظة:', self.prov)
        self.action = QComboBox()
        self.action.addItems(['تأشير ✔ أجابت', 'إلغاء ✘ لم تُجب'])
        f.addRow('الإجراء:', self.action)
        lay.addLayout(f)

        warn = QLabel('⚠️ ملاحظة: سيتم تطبيق العملية على كل الملفات الظاهرة حالياً بعد البحث أو التصفية أو نطاق رقم الملف.')
        warn.setStyleSheet('color:#b52a1c;font-size:12px')
        warn.setWordWrap(True)
        lay.addWidget(warn)

        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText('✅ تنفيذ')
        bb.button(QDialogButtonBox.Cancel).setText('إلغاء')
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)

    def get_data(self):
        return self.prov.currentIndex(), self.action.currentIndex()


class Main(QWidget):
    def __init__(self):
        super().__init__()
        self.db = DB()
        self.rows = {}
        self._hay = {}
        self.shown = []
        self.rowOf = {}
        self.tab = 'all'
        self.sortCol = -1
        self.sortAsc = True
        self.undo_st = []
        self.redo_st = []
        self.big = QFont('Segoe UI Symbol', 17, QFont.Bold)
        self.setWindowTitle(f'أرشيف التحقيق العقاري — {ORG}')
        self.settings = QSettings('DCCF', 'ArchiveTahqiqPro')
        self.dark_mode = self.settings.value('dark_mode', False, type=bool)
        self.resize(1600, 920)

        root = QVBoxLayout(self)
        root.setContentsMargins(14, 8, 14, 6)
        root.setSpacing(8)

        hdr = QFrame(objectName='hdr')
        hdr.setFixedHeight(86)
        h = QHBoxLayout(hdr)
        h.setContentsMargins(22, 4, 22, 4)

        tx = QVBoxLayout()
        tx.setSpacing(0)
        tx.addWidget(QLabel('أرشيف التحقيق العقاري', objectName='title'))
        tx.addWidget(QLabel('سجل متابعة داخلي · ولاية الجلفة  —  ' + ' · '.join(PROV), objectName='gold'))
        h.addLayout(tx)
        h.addStretch()

        self.themeBtn = QPushButton(objectName='hdrBtn')
        self.themeBtn.setCursor(Qt.PointingHandCursor)
        self.themeBtn.setToolTip('تبديل الوضع النهاري / الليلي')
        self.themeBtn.clicked.connect(self.toggle_theme)
        h.addWidget(self.themeBtn)

        for t, fn in [('📂 استرجاع نسخة', self.restore_backup), ('🔑 تغيير كلمة المرور', lambda: ChangePwDlg(self).exec_()), ('🔒 قفل البرنامج', self.lock_app)]:
            b = QPushButton(t, objectName='hdrBtn')
            b.setCursor(Qt.PointingHandCursor)
            b.clicked.connect(lambda _, f=fn: f())
            h.addWidget(b)
        h.addSpacing(10)
        h.addWidget(QLabel(ORG, objectName='gold'))
        root.addWidget(hdr)

        st = QHBoxLayout()
        self.stat = {}
        for k, t in [('total', 'إجمالي الملفات'), ('done', 'مكمل الإجابة ✔'), ('pend', 'معلق لم يكتمل'), ('noatt', 'بلا مرفق'), ('pct', 'نسبة الإنجاز')]:
            c = QFrame(objectName='card')
            c.setFixedHeight(58)
            v = QHBoxLayout(c)
            v.setContentsMargins(14, 2, 14, 2)
            v.addWidget(QLabel(t))
            v.addStretch()
            self.stat[k] = QLabel('0', objectName='big')
            v.addWidget(self.stat[k])
            st.addWidget(c)
        root.addLayout(st)

        pb = QHBoxLayout()
        self.bars = []
        for _ in PROV:
            b = QProgressBar()
            pb.addWidget(b)
            self.bars.append(b)
        root.addLayout(pb)

        self.timer = QTimer(self)
        self.timer.setSingleShot(True)
        self.timer.setInterval(250)
        self.timer.timeout.connect(self.refresh)

        sr = QHBoxLayout()
        self.search = QLineEdit(objectName='search', placeholderText='🔍  ابحث... يمكنك كتابة عدة كلمات معاً  (Ctrl+F)')
        self.search.setClearButtonEnabled(True)
        self.search.textChanged.connect(lambda _: self.timer.start())
        self.scope = combo(['البحث في كل الحقول', 'رقم الملف فقط', 'الجهة المرسلة فقط', 'رقم الإرسال فقط', 'الملاحظات فقط'])
        self.scope.currentIndexChanged.connect(self.refresh)
        sr.addWidget(self.search, 4)
        sr.addWidget(self.scope, 1)
        root.addLayout(sr)

        f1 = QHBoxLayout()
        self.fStatus = combo(['كل الحالات', 'مكمل الإجابة ✔', 'معلق لم يكتمل'])
        self.fSender = QComboBox()
        self.fYear = QComboBox()
        self.fProv = combo(['كل المحافظات'] + [f'✘ لم تُجب: {n}' for n in PROV] + [f'✔ أجابت: {n}' for n in PROV])
        self.fCount = combo(['كل أعداد الردود'] + [f'{i} من 5 ردود' for i in range(6)])
        self.fAtt = combo(['المرفق: الكل', 'بمرفق 📎', 'بلا مرفق'])
        for c in (self.fStatus, self.fSender, self.fProv, self.fCount, self.fAtt, self.fYear):
            f1.addWidget(c, 1)
            c.currentIndexChanged.connect(self.refresh)
        root.addLayout(f1)

        f2 = QHBoxLayout()
        self.nFrom = QLineEdit(placeholderText='مثال: 800 أو 2025/800')
        self.nTo = QLineEdit(placeholderText='مثال: 900 أو 2025/900')
        for w in (self.nFrom, self.nTo):
            w.setMaximumWidth(200)
            w.setClearButtonEnabled(True)
            w.textChanged.connect(lambda _: self.timer.start())
            w.setToolTip('اكتب الرقم التسلسلي مثل 800 إلى 900 أو الرقم الكامل مثل 2025/800 إلى 2025/900')
        self.fSort = combo([
            'الترتيب: حسب الإدخال', 'الأقدم إرسالاً أولاً', 'الأحدث إرسالاً أولاً',
            'الأكثر نقصاً في الردود', 'الأقرب إلى الاكتمال', 'رقم الملف تصاعدياً',
            'رقم الملف تنازلياً', 'الجهة المرسلة (أ ← ي)'
        ])
        self.fSort.currentIndexChanged.connect(self.sort_combo)
        rst = QPushButton('↺  مسح كل الفلاتر')
        rst.clicked.connect(self.reset)
        self.bulkBtn = QPushButton('⚡ إجراء جماعي على النتائج')
        self.bulkBtn.setCursor(Qt.PointingHandCursor)
        self.bulkBtn.setStyleSheet('font-weight:bold;color:#123b2e')
        self.bulkBtn.clicked.connect(self.do_bulk_action)
        self.count = QLabel(objectName='hint')
        for w in (QLabel('📁 رقم الملف:  من'), self.nFrom, QLabel('إلى'), self.nTo):
            f2.addWidget(w)
        f2.addSpacing(20)
        f2.addWidget(self.fSort, 1)
        f2.addWidget(rst)
        f2.addWidget(self.bulkBtn)
        f2.addStretch()
        f2.addWidget(self.count)
        root.addLayout(f2)

        tr = QHBoxLayout()
        self.tabs = {}
        self.grp = QButtonGroup(self)
        for k, t in [('all', 'سجل العمل'), ('pend', 'المعلقة'), ('done', 'المكملة'), ('noatt', 'بلا مرفق')]:
            b = QPushButton(t, objectName='tab', checkable=True)
            b.setChecked(k == 'all')
            self.grp.addButton(b)
            b.clicked.connect(lambda _, k=k: self.set_tab(k))
            tr.addWidget(b)
            self.tabs[k] = (b, t)
        tr.addSpacing(12)
        self.bUndo = QPushButton('↶', objectName='hist')
        self.bUndo.clicked.connect(self.undo)
        self.bRedo = QPushButton('↷', objectName='hist')
        self.bRedo.clicked.connect(self.redo)
        for b in (self.bUndo, self.bRedo):
            b.setMinimumHeight(44)
            b.setCursor(Qt.PointingHandCursor)
            tr.addWidget(b)
        self.msg = QLabel(objectName='msg')
        self.msg.hide()
        tr.addWidget(self.msg)
        tr.addStretch()
        for t, fn in [('✏ تعديل', self.edit_sel), ('🗑 حذف', self.delete), ('📎 المرفق', self.open_att), ('📥 استيراد إكسل', self.import_xlsx), ('📊 تصدير إكسل', self.export_xlsx), ('📄 تصدير PDF', self.export_pdf), ('💾 نسخة احتياطية', self.backup)]:
            b = QPushButton(t)
            b.clicked.connect(lambda _, f=fn: f())
            tr.addWidget(b)
        add = QPushButton('➕   إضافة ملف جديد', objectName='addBtn')
        add.setMinimumHeight(54)
        add.setCursor(Qt.PointingHandCursor)
        add.setToolTip('Alt+N')
        add.clicked.connect(lambda: self.edit())
        tr.addWidget(add)
        root.addLayout(tr)

        self.tbl = QTableWidget(0, len(COLS))
        self.tbl.setHorizontalHeaderLabels(COLS)
        self.tbl.setAlternatingRowColors(True)
        self.tbl.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tbl.setSelectionMode(QAbstractItemView.SingleSelection)
        self.tbl.setWordWrap(False)
        self.tbl.setMinimumHeight(380)
        vh = self.tbl.verticalHeader(); vh.hide(); vh.setSectionResizeMode(QHeaderView.Fixed); vh.setDefaultSectionSize(48)
        hh = self.tbl.horizontalHeader(); hh.setSectionResizeMode(QHeaderView.Stretch); hh.setMinimumHeight(44)
        for c, wd in [(0, 60), (C_STATUS, 190), (C_ATT, 80)] + [(c, 105) for c in range(C_PROV, C_PROV + 5)]:
            hh.setSectionResizeMode(c, QHeaderView.Fixed)
            self.tbl.setColumnWidth(c, wd)
        hh.setSectionsClickable(True)
        hh.setHighlightSections(False)
        hh.sectionClicked.connect(self.head_sort)
        self.tbl.cellClicked.connect(self.clicked)
        self.tbl.cellDoubleClicked.connect(self.dbl)
        self.tbl.setContextMenuPolicy(Qt.CustomContextMenu)
        self.tbl.customContextMenuRequested.connect(self.menu)
        root.addWidget(self.tbl, 10)

        root.addWidget(QLabel('✔ أجابت  ·  ✘ لم تُجب  ·  انقر على الخانة لتبديل الحالة  ·  ↶ Ctrl+Z  ·  ↷ Ctrl+Y  ·  🔒 Ctrl+L  ·  الزر الأيمن لخيارات إضافية  ·  ' + ORG + ' © 2025', objectName='hint'))

        fb = QHBoxLayout()
        fb.addStretch()
        self.fab = QPushButton(objectName='fab')
        self.fab.clicked.connect(self.all_files)
        fb.addWidget(self.fab)
        fb.addStretch()
        root.addLayout(fb)

        for k, fn in [('Alt+N', lambda: self.edit()), ('Ctrl+F', lambda: (self.search.setFocus(), self.search.selectAll())), ('F5', lambda: (self.reload(), self.refresh())), ('Ctrl+Z', self.undo), ('Ctrl+Y', self.redo), ('Ctrl+Shift+Z', self.redo), ('Ctrl+L', self.lock_app)]:
            QShortcut(QKeySequence(k), self, activated=fn)
        QShortcut(QKeySequence(Qt.Key_Delete), self.tbl, activated=self.delete)

        self.apply_theme(False)
        self.reload(); self.upd_hist(); self.refresh()

    def reload(self):
        self.rows = {r['id']: r for r in self.db.all()}
        self._hay.clear()

    def set_row(self, i, row):
        self._hay.pop(i, None)
        if row:
            self.rows[i] = row
        else:
            self.rows.pop(i, None)

    def hay(self, r, sc=0):
        h = self._hay.get(r['id'])
        if h is None:
            full = norm(' '.join(str(x) for x in (r['file_no'], r['sender'], r['send_no'], r['send_date'], iso2d(r['send_date']), r['notes'], 'مكمل' if done(r) == 5 else 'معلق')))
            h = (full, norm(r['file_no']), norm(r['sender']), norm(r['send_no']), norm(r['notes']))
            self._hay[r['id']] = h
        return h[sc]

    def lock_app(self):
        app = QApplication.instance()
        app.setQuitOnLastWindowClosed(False)
        self.hide()
        if LoginDlg().exec_() == QDialog.Accepted:
            self.showMaximized(); app.setQuitOnLastWindowClosed(True)
        else:
            app.quit()

    def restore_backup(self):
        p, _ = QFileDialog.getOpenFileName(self, 'اختر ملف النسخة الاحتياطية', os.path.expanduser('~'), 'نسخة احتياطية (*.zip)')
        if not p:
            return
        try:
            with zipfile.ZipFile(p) as z:
                if 'archive.db' not in z.namelist():
                    return QMessageBox.warning(self, 'خطأ', 'هذا الملف ليس نسخة احتياطية للبرنامج')
                tmp = tempfile.mkdtemp(); z.extractall(tmp)
        except zipfile.BadZipFile:
            return QMessageBox.warning(self, 'خطأ', 'ملف النسخة تالف أو غير صالح')

        try:
            c = sqlite3.connect(os.path.join(tmp, 'archive.db'))
            n = c.execute('SELECT COUNT(*) FROM files').fetchone()[0]
            c.close()
        except sqlite3.Error:
            shutil.rmtree(tmp, ignore_errors=True)
            return QMessageBox.warning(self, 'خطأ', 'قاعدة البيانات داخل النسخة غير صالحة')

        if QMessageBox.question(self, 'تأكيد الاسترجاع', f'النسخة تحتوي على {n} ملف.\nالبيانات الحالية ({len(self.rows)} ملف) سيتم استبدالها.\n\nسيتم حفظ نسخة أمان من بياناتك الحالية في مجلد Documents.\n\nمتابعة؟') != QMessageBox.Yes:
            shutil.rmtree(tmp, ignore_errors=True)
            return

        docs = os.path.join(os.path.expanduser('~'), 'Documents'); os.makedirs(docs, exist_ok=True)
        safe = shutil.make_archive(os.path.join(docs, f'ارشيف_قبل_الاسترجاع_{time.strftime("%Y%m%d_%H%M%S")}'), 'zip', DATA_DIR)
        try:
            self.db.c.close()
            shutil.copy2(os.path.join(tmp, 'archive.db'), os.path.join(DATA_DIR, 'archive.db'))
            src = os.path.join(tmp, 'attachments')
            if os.path.isdir(src):
                shutil.copytree(src, ATT_DIR, dirs_exist_ok=True)
        except Exception as e:
            QMessageBox.warning(self, 'خطأ', f'تعذّر الاسترجاع:\n{e}\n\nنسخة الأمان موجودة في:\n{safe}')
        finally:
            self.db = DB(); shutil.rmtree(tmp, ignore_errors=True)
        self.reload(); self.undo_st.clear(); self.redo_st.clear(); self.upd_hist(); self.refresh()
        QMessageBox.information(self, 'تم', f'تم استرجاع النسخة الاحتياطية ✓\nعدد الملفات: {len(self.rows)}\n\nنسخة أمان:\n{safe}')

    def record(self, label, ch):
        ch = [c for c in ch if c[1] != c[2]]
        if ch:
            self.undo_st.append((label, ch)); self.redo_st.clear(); del self.undo_st[:-MAX_HIST]
        self.upd_hist()

    def apply(self, ch, undo):
        tgt = [(i, b if undo else a) for i, b, a in ch]
        try:
            for i, _ in tgt:
                self.db.raw_delete(i)
            for _, row in tgt:
                if row:
                    self.db.put(row)
            self.db.c.commit()
        except sqlite3.Error as e:
            self.db.c.rollback(); QMessageBox.warning(self, 'خطأ', f'تعذّر تنفيذ العملية:\n{e}'); return False
        for i, row in tgt:
            self.set_row(i, row)
        return True

    def after(self, ch):
        if len(ch) == 1:
            self.row_changed(ch[0][0])
        else:
            self.refresh()

    def undo(self):
        if not self.undo_st:
            return
        lbl, ch = self.undo_st.pop()
        if self.apply(ch, True):
            self.redo_st.append((lbl, ch)); self.flash(f'↶ تم التراجع عن: {lbl}'); self.after(ch)
        else:
            self.undo_st.append((lbl, ch))
        self.upd_hist()

    def redo(self):
        if not self.redo_st:
            return
        lbl, ch = self.redo_st.pop()
        if self.apply(ch, False):
            self.undo_st.append((lbl, ch)); self.flash(f'↷ تمت إعادة: {lbl}'); self.after(ch)
        else:
            self.redo_st.append((lbl, ch))
        self.upd_hist()

    def upd_hist(self):
        self.bUndo.setEnabled(bool(self.undo_st)); self.bRedo.setEnabled(bool(self.redo_st))
        self.bUndo.setToolTip(f'تراجع عن: {self.undo_st[-1][0]}   (Ctrl+Z)' if self.undo_st else 'لا توجد عملية للتراجع عنها')
        self.bRedo.setToolTip(f'إعادة: {self.redo_st[-1][0]}   (Ctrl+Y)' if self.redo_st else 'لا توجد عملية لإعادتها')

    def flash(self, text):
        self.msg.setText(text); self.msg.show(); QTimer.singleShot(4000, lambda: self.msg.text() == text and self.msg.hide())

    def do_bulk_action(self):
        if not self.shown:
            return QMessageBox.information(self, 'تنبيه', 'لا توجد ملفات معروضة لتطبيق الإجراء عليها.\nحدد نطاقاً أو استعمل البحث/الفلاتر أولاً.')
        dlg = BulkActionDlg(self, len(self.shown))
        if dlg.exec_() != QDialog.Accepted:
            return
        prov_idx, action_idx = dlg.get_data()
        is_check = action_idx == 0
        today = date.today().isoformat(); changes = []
        for r in list(self.shown):
            i = r['id']; b = self.rows[i]; a = dict(b); changed = False
            if prov_idx == 0:
                for k in range(5):
                    new_val = '1' if is_check else ''
                    if a[f'p{k}'] != new_val:
                        a[f'p{k}'] = new_val; a[f't{k}'] = today if is_check else ''; changed = True
            else:
                k = prov_idx - 1; new_val = '1' if is_check else ''
                if a[f'p{k}'] != new_val:
                    a[f'p{k}'] = new_val; a[f't{k}'] = today if is_check else ''; changed = True
            if changed:
                changes.append((i, b, a))
        if not changes:
            return QMessageBox.information(self, 'تنبيه', 'لا توجد تغييرات: الملفات المعروضة تحمل نفس الحالة مسبقاً.')
        prov_name = 'كل المحافظات' if prov_idx == 0 else PROV[prov_idx - 1]
        action_name = 'تأشير ✔' if is_check else 'إلغاء ✘'
        if QMessageBox.question(self, 'تأكيد الإجراء الجماعي', f'سيتم تنفيذ: {action_name}\nالمحافظة: {prov_name}\nعدد الملفات المتأثرة: {len(changes)}\n\nهل تريد المتابعة؟') != QMessageBox.Yes:
            return
        if self.apply(changes, undo=False):
            self.record(f'{action_name} {prov_name} لـ {len(changes)} ملف', changes)
            self.refresh()
            QMessageBox.information(self, 'تم', f'تم تطبيق الإجراء على {len(changes)} ملف بنجاح ✓\nيمكنك التراجع عن العملية بالسهم ↶ أو Ctrl+Z')

    def flt(self):
        return {'toks': norm(self.search.text()).split(), 'sc': self.scope.currentIndex(), 'tab': self.tab, 'stt': self.fStatus.currentIndex(), 'pv': self.fProv.currentIndex(), 'cn': self.fCount.currentIndex(), 'at': self.fAtt.currentIndex(), 'snd': self.fSender.currentText() if self.fSender.currentIndex() > 0 else None, 'yr': self.fYear.currentText() if self.fYear.currentIndex() > 0 else None, 'lo': self.nFrom.text().strip(), 'hi': self.nTo.text().strip()}

    def match(self, r, F):
        n = done(r)
        tab = F['tab']; pv = F['pv']
        if (tab == 'pend' and n == 5) or (tab == 'done' and n < 5) or (tab == 'noatt' and r['att']):
            return False
        if (F['stt'] == 1 and n < 5) or (F['stt'] == 2 and n == 5):
            return False
        if F['snd'] and r['sender'] != F['snd']:
            return False
        if 1 <= pv <= 5 and r[f'p{pv - 1}']:
            return False
        if pv >= 6 and not r[f'p{pv - 6}']:
            return False
        if F['cn'] > 0 and n != F['cn'] - 1:
            return False
        if (F['at'] == 1 and not r['att']) or (F['at'] == 2 and r['att']):
            return False
        if F['yr'] and not r['send_date'].startswith(F['yr']):
            return False
        if (F['lo'] or F['hi']) and not in_range(r['file_no'], F['lo'], F['hi']):
            return False
        if F['toks']:
            h = self.hay(r, F['sc'])
            if not all(t in h for t in F['toks']):
                return False
        return True

    def cell(self, i, j, t, bg=None, fg=None):
        it = QTableWidgetItem(t)
        it.setTextAlignment(Qt.AlignCenter)
        it.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
        if bg is not None:
            it.setBackground(bg)
        if fg is not None:
            it.setForeground(fg)
        self.tbl.setItem(i, j, it)
        return it

    def fill(self, cb, first, items):
        if [cb.itemText(k) for k in range(1, cb.count())] == items:
            return
        cur = cb.currentText(); cb.blockSignals(True); cb.clear(); cb.addItem(first); cb.addItems(items); cb.setCurrentIndex(max(cb.findText(cur), 0)); cb.blockSignals(False)

    def fill_row(self, i, r, toks):
        vals = [str(i + 1), r['file_no'], r['sender'], r['send_no'], iso2d(r['send_date'])]
        for j, v in enumerate(vals):
            it = self.cell(i, j, v)
            if toks and 1 <= j <= 3 and any(t in self.hay(r, j) for t in toks):
                it.setBackground(HL)
        self.tbl.item(i, 0).setData(Qt.UserRole, r['id'])
        for k in range(5):
            ok = bool(r[f'p{k}'])
            it = self.cell(i, C_PROV + k, YES if ok else NO, GREEN_BG if ok else RED_BG, GREEN_FG if ok else RED_FG)
            it.setFont(self.big)
            if ok:
                td = iso2d(r.get(f't{k}', ''))
                it.setToolTip(f'{PROV[k]}: أجابت ✔  ({td})')
            else:
                it.setToolTip(f'{PROV[k]}: لم تُجب ✘ — انقر للتأشير')
        n = done(r)
        if n == 5:
            self.cell(i, C_STATUS, 'مكمل الإجابة ✔', DONE_BG, DONE_FG)
        else:
            self.cell(i, C_STATUS, f'معلق لم يكتمل  {n}/5', PEND_BG, PEND_FG)
        self.cell(i, C_ATT, '📎' if r['att'] else '—')

    def update_stats(self):
        tot = len(self.rows)
        cnt5 = [0] * 5; dn = 0; na = 0
        for r in self.rows.values():
            n = 0
            for k in range(5):
                if r[f'p{k}']:
                    cnt5[k] += 1; n += 1
            if n == 5: dn += 1
            if not r['att']: na += 1
        for k, v in [('total', tot), ('done', dn), ('pend', tot - dn), ('noatt', na), ('pct', f'{round(dn / tot * 100) if tot else 0}%')]:
            self.stat[k].setText(str(v))
        for k, b in enumerate(self.bars):
            b.setMaximum(max(tot, 1)); b.setValue(cnt5[k]); b.setFormat(f'{PROV[k]}   {cnt5[k]}/{tot}')
        cnt = {'all': tot, 'pend': tot - dn, 'done': dn, 'noatt': na}
        for k, (b, t) in self.tabs.items():
            b.setText(f'{t}  ({cnt[k]})')
        self.fab.setText(f'🗂   كل الملفات   ( {tot} )')

    def refresh(self, *_):
        keep = self.cur()
        rows = sorted(self.rows.values(), key=lambda r: r['id'])
        self.update_stats()
        self.fill(self.fSender, 'كل الجهات', sorted({r['sender'] for r in rows if r['sender']}, key=norm))
        self.fill(self.fYear, 'كل السنوات', sorted({r['send_date'][:4] for r in rows if r['send_date']}, reverse=True))
        F = self.flt(); L = [r for r in rows if self.match(r, F)]
        if self.sortCol >= 0:
            c = self.sortCol
            def key(r):
                if c in (1, 3):
                    return nat(r['file_no'] if c == 1 else r['send_no'])
                if c == 2: return norm(r['sender'])
                if c == 4: return r['send_date']
                if C_PROV <= c < C_PROV + 5: return bool(r[f'p{c - C_PROV}'])
                if c == C_STATUS: return done(r)
                if c == C_ATT: return bool(r['att'])
                return r['id']
            try:
                L.sort(key=key, reverse=not self.sortAsc)
            except TypeError:
                L.sort(key=lambda r: str(key(r)), reverse=not self.sortAsc)
        else:
            so = self.fSort.currentIndex()
            if so == 1: L.sort(key=lambda r: r['send_date'] or '9999')
            elif so == 2: L.sort(key=lambda r: r['send_date'], reverse=True)
            elif so == 3: L.sort(key=lambda r: (done(r), r['send_date'] or '9999'))
            elif so == 4: L.sort(key=lambda r: (-done(r), r['send_date'] or '9999'))
            elif so == 5: L.sort(key=lambda r: nat(r['file_no']))
            elif so == 6: L.sort(key=lambda r: nat(r['file_no']), reverse=True)
            elif so == 7: L.sort(key=lambda r: norm(r['sender']))
        self.shown = L; self.rowOf = {r['id']: n for n, r in enumerate(L)}
        t = self.tbl; t.setUpdatesEnabled(False)
        try:
            t.clearContents(); t.setRowCount(len(L))
            for n, r in enumerate(L):
                self.fill_row(n, r, F['toks'])
            if keep in self.rowOf: t.selectRow(self.rowOf[keep])
        finally:
            t.setUpdatesEnabled(True)
        self.count.setText(f'🔎 النتائج: {len(L)} من {len(rows)}')
        if hasattr(self, 'bulkBtn'):
            self.bulkBtn.setEnabled(bool(L))

    def row_changed(self, i):
        self.update_stats()
        r = self.rows.get(i); tr = self.rowOf.get(i)
        if tr is None:
            if r is not None: self.refresh(); return
        F = self.flt()
        if r is not None and self.match(r, F):
            self.shown[tr] = r; self.fill_row(tr, r, F['toks'])
        else:
            self.tbl.removeRow(tr); del self.shown[tr]; self.rowOf = {x['id']: n for n, x in enumerate(self.shown)}
            for n in range(tr, len(self.shown)): self.tbl.item(n, 0).setText(str(n + 1))
        self.count.setText(f'🔎 النتائج: {len(self.shown)} من {len(self.rows)}')

    def head_sort(self, c):
        if self.sortCol != c:
            self.sortCol = c; self.sortAsc = True
        elif self.sortAsc:
            self.sortAsc = False
        else:
            self.sortCol = -1
        hh = self.tbl.horizontalHeader(); hh.setSortIndicatorShown(self.sortCol >= 0)
        if self.sortCol >= 0:
            hh.setSortIndicator(c, Qt.AscendingOrder if self.sortAsc else Qt.DescendingOrder)
        self.refresh()

    def sort_combo(self, *_):
        self.sortCol = -1; self.tbl.horizontalHeader().setSortIndicatorShown(False); self.refresh()

    def set_tab(self, k):
        self.tab = k; self.refresh()

    def reset(self):
        for c in (self.fStatus, self.fSender, self.fProv, self.fCount, self.fAtt, self.fYear, self.fSort, self.scope):
            c.blockSignals(True); c.setCurrentIndex(0); c.blockSignals(False)
        for w in (self.search, self.nFrom, self.nTo):
            w.blockSignals(True); w.clear(); w.blockSignals(False)
        self.sortCol = -1; self.tbl.horizontalHeader().setSortIndicatorShown(False); self.refresh()

    def cur(self):
        r = self.tbl.currentRow(); return self.tbl.item(r, 0).data(Qt.UserRole) if r >= 0 and self.tbl.item(r, 0) else None

    def done_msg(self, no):
        QMessageBox.information(self, 'اكتمل الملف', f'🎉 اكتمل الملف {no} — أجابت المحافظات الخمس')

    def clicked(self, row, col):
        if C_PROV <= col < C_PROV + 5:
            i = self.tbl.item(row, 0).data(Qt.UserRole); now = time.time()
            if self._last and self._last[:2] == (i, col) and now - self._last[2] < 0.5:
                return
            self._last = (i, col, now)
            self.toggle(i, col - C_PROV)

    def toggle(self, i, k):
        b = self.rows[i]; on = not b[f'p{k}']
        self.db.save({f'p{k}': '1' if on else '', f't{k}': date.today().isoformat() if on else ''}, i)
        a = self.db.get(i); self.set_row(i, a)
        self.record(f'{"تأشير ✔" if on else "إلغاء ✘"} {PROV[k]} — الملف {a["file_no"]}', [(i, b, a)])
        self.row_changed(i)
        if on and done(a) == 5: self.done_msg(a['file_no'])

    def set_all(self, i, v):
        b = self.rows[i]; today = date.today().isoformat(); d = {}
        for k in range(5):
            d[f'p{k}'] = (b[f'p{k}'] or '1') if v else ''
            d[f't{k}'] = (b[f't{k}'] or today) if v else ''
        self.db.save(d, i)
        a = self.db.get(i); self.set_row(i, a)
        self.record(f'{"تأشير كل المحافظات" if v else "إلغاء كل الردود"} — الملف {a["file_no"]}', [(i, b, a)])
        self.row_changed(i)
        if v and done(b) < 5: self.done_msg(a['file_no'])

    def dbl(self, row, col):
        if C_PROV <= col < C_PROV + 5:
            return
        if col == C_ATT: return self.open_att()
        self.edit(self.tbl.item(row, 0).data(Qt.UserRole))

    def menu(self, pos):
        row = self.tbl.rowAt(pos.y())
        if row < 0: return
        self.tbl.selectRow(row); i = self.cur(); m = QMenu(self)
        m.addAction('✏  تعديل الملف', lambda: self.edit(i))
        m.addAction('✔  أجابت كل المحافظات', lambda: self.set_all(i, True))
        m.addAction('✘  إلغاء كل الردود', lambda: self.set_all(i, False))
        m.addSeparator(); m.addAction('📎  فتح المرفق', self.open_att); m.addAction('🗑  حذف الملف', self.delete)
        m.exec_(self.tbl.viewport().mapToGlobal(pos))

    def edit(self, i=None):
        b = self.rows.get(i) if i else None
        senders = sorted({x['sender'] for x in self.rows.values() if x['sender']} | set(SENDERS), key=norm)
        d = FileDlg(self, senders, b)
        while d.exec_():
            try:
                nid = self.db.save(d.data(), i)
                break
            except sqlite3.IntegrityError:
                QMessageBox.warning(self, 'تنبيه', 'رقم الملف موجود مسبقاً في الأرشيف')
        else:
            return
        a = self.db.get(nid); self.set_row(nid, a)
        no = a['file_no']
        self.record(f'تعديل الملف {no}' if b else f'إضافة الملف {no}', [(nid, b, a)])
        if done(a) == 5 and (not b or done(b) < 5): self.done_msg(no)
        self.refresh()

    def edit_sel(self):
        i = self.cur()
        if i: self.edit(i)
        else: QMessageBox.information(self, 'تنبيه', 'اختر ملفاً من الجدول أولاً')

    def delete(self):
        i = self.cur();
        if not i: return QMessageBox.information(self, 'تنبيه', 'اختر ملفاً من الجدول أولاً')
        b = self.rows[i]
        if QMessageBox.question(self, 'تأكيد الحذف', f'حذف الملف {b["file_no"]}؟\n(يمكنك التراجع بالسهم ↶)') != QMessageBox.Yes:
            return
        self.db.delete(i); self.set_row(i, None); self.record(f'حذف الملف {b["file_no"]}', [(i, b, None)]); self.row_changed(i)

    def open_att(self, i=None):
        i = i or self.cur(); r = self.rows.get(i) if i else None
        if not r or not r['att']: return QMessageBox.information(self, 'المرفق', 'لا يوجد مرفق — أضفه من نافذة التعديل')
        p = os.path.join(ATT_DIR, r['att'])
        if os.path.exists(p): os.startfile(p)
        else: QMessageBox.warning(self, 'خطأ', 'ملف المرفق غير موجود')

    def all_files(self):
        d = QDialog(self, Qt.Window)
        d.setWindowTitle(f'كل الملفات المضافة — {ORG}')
        d.resize(1450, 820)
        lay = QVBoxLayout(d)
        top = QHBoxLayout(); sr = QLineEdit(objectName='search', placeholderText='🔍  بحث في كل البيانات... (عدة كلمات معاً)'); sr.setClearButtonEnabled(True); cnt = QLabel(objectName='hint'); xb = QPushButton('📊 تصدير هذه القائمة إلى إكسل'); top.addWidget(sr, 1); top.addWidget(cnt); top.addWidget(xb); lay.addLayout(top)
        cols = ['العدد', 'رقم الملف', 'الجهة المرسلة', 'رقم الإرسال', 'تاريخ الإرسال', *PROV, 'الحالة', 'المرفق', 'ملاحظات']
        t = QTableWidget(0, len(cols)); t.setHorizontalHeaderLabels(cols); t.setAlternatingRowColors(True); t.setSelectionBehavior(QAbstractItemView.SelectRows); t.setSelectionMode(QAbstractItemView.SingleSelection); t.setWordWrap(False)
        vh = t.verticalHeader(); vh.hide(); vh.setSectionResizeMode(QHeaderView.Fixed); vh.setDefaultSectionSize(46)
        hh = t.horizontalHeader(); hh.setSectionResizeMode(QHeaderView.Interactive); hh.setStretchLastSection(True); hh.setMinimumHeight(44); hh.setSortIndicator(0, Qt.AscendingOrder)
        lay.addWidget(t, 1)
        bot = QHBoxLayout(); bot.addWidget(QLabel('انقر مرتين على أي ملف لتعديله  ·  انقر على رأس العمود للترتيب  ·  انقر مرتين على 📎 لفتح المرفق', objectName='hint')); bot.addStretch(); cb = QPushButton('إغلاق'); cb.clicked.connect(d.accept); bot.addWidget(cb); lay.addLayout(bot)
        big = QFont('Segoe UI Symbol', 16, QFont.Bold); cur = []; first = [True]
        def fill():
            rows = sorted(self.rows.values(), key=lambda r: r['id']); toks = norm(sr.text()).split(); L = [r for r in rows if all(x in self.hay(r, 0) + ' ' + norm(r['att']) for x in toks)]; cur[:] = L; t.setUpdatesEnabled(False); t.setSortingEnabled(False); t.clearContents(); t.setRowCount(len(L))
            for i, r in enumerate(L):
                n = done(r); vals = [(i+1, i+1), (r['file_no'], nat(r['file_no'])), (r['sender'], norm(r['sender'])), (r['send_no'], nat(r['send_no'])), (iso2d(r['send_date']), r['send_date'])]
                for j, (v, k) in enumerate(vals): t.setItem(i, j, SItem(v, k)); t.item(i, 0).setData(Qt.UserRole, r['id'])
                for k in range(5):
                    ok = bool(r[f'p{k}']); it = SItem(YES if ok else NO, int(ok)); it.setBackground(GREEN_BG if ok else RED_BG); it.setForeground(GREEN_FG if ok else RED_FG); it.setFont(big); t.setItem(i, 5+k, it)
                it = SItem('مكمل الإجابة ✔' if n == 5 else f'معلق لم يكتمل  {n}/5', n); it.setBackground(DONE_BG if n == 5 else PEND_BG); it.setForeground(DONE_FG if n == 5 else PEND_FG); t.setItem(i, 10, it)
                t.setItem(i, 11, SItem(('📎 ' + r['att'].split('_', 1)[-1]) if r['att'] else '—'))
                nt = SItem(r['notes'] or '—'); nt.setTextAlignment(Qt.AlignVCenter | Qt.AlignRight); t.setItem(i, 12, nt)
            t.setSortingEnabled(True)
            if first[0]: t.resizeColumnsToContents(); first[0] = False
            t.setUpdatesEnabled(True); cnt.setText(f'🔎 {len(L)} من {len(rows)} ملف')
        def dbl_row(row, col):
            i = t.item(row, 0).data(Qt.UserRole)
            if col == 11: return self.open_att(i)
            self.edit(i); fill()
        tm = QTimer(d); tm.setSingleShot(True); tm.setInterval(250); tm.timeout.connect(fill)
        sr.textChanged.connect(lambda _: tm.start()); t.cellDoubleClicked.connect(dbl_row); xb.clicked.connect(lambda: self.export_xlsx(list(cur)))
        fill(); d.exec_()

    def export_xlsx(self, L=None):
        L = self.shown if L is None else L
        p, _ = QFileDialog.getSaveFileName(self, 'تصدير إكسل', f'ارشيف_التحقيق_{date.today()}.xlsx', 'Excel (*.xlsx)')
        if not p: return
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = 'الأرشيف'; ws.sheet_view.rightToLeft = True; ws.append(HEAD + ['ملاحظات'])
        for i, r in enumerate(L):
            ws.append([i+1, r['file_no'], r['sender'], r['send_no'], iso2d(r['send_date']), *[YES if r[f'p{k}'] else NO for k in range(5)], 'مكمل الإجابة' if done(r) == 5 else f'معلق لم يكتمل {done(r)}/5', r['notes']])
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF', size=12); c.fill = PatternFill('solid', fgColor='123B2E')
        green = PatternFill('solid', fgColor='DFF0E5'); red = PatternFill('solid', fgColor='FBE9E8'); fg = Font(bold=True, size=14, color='136B40'); fr = Font(bold=True, size=14, color='B52A1C'); center = Alignment(horizontal='center', vertical='center')
        for row in ws.iter_rows():
            for c in row: c.alignment = center
        for row in ws.iter_rows(min_row=2):
            for c in row[5:10]:
                ok = c.value == YES; c.fill = green if ok else red; c.font = fg if ok else fr
        for col, w in zip('ABCDEFGHIJKL', [7, 14, 28, 14, 14, 12, 12, 12, 12, 12, 20, 30]): ws.column_dimensions[col].width = w
        wb.save(p)
        QMessageBox.information(self, 'تم', f'تم تصدير {len(L)} ملف إلى إكسل ✓\n{p}')

    def import_xlsx(self):
        p, _ = QFileDialog.getOpenFileName(self, 'استيراد إكسل', '', 'Excel (*.xlsx)')
        if not p: return
        from openpyxl import load_workbook
        try:
            rows = [list(r) for r in load_workbook(p, data_only=True, read_only=True).active.iter_rows(values_only=True)]
        except Exception as e:
            return QMessageBox.warning(self, 'خطأ', f'تعذّر فتح ملف الإكسل:\n{e}')
        h = next((i for i, r in enumerate(rows) if 'رقم الملف' in [str(x or '').strip() for x in r]), None)
        if h is None: return QMessageBox.warning(self, 'خطأ', 'لم يتم العثور على عمود "رقم الملف"')
        head = [str(x or '').strip() for x in rows[h]]
        raw = lambda r, n: r[head.index(n)] if n in head and head.index(n) < len(r) else None
        g = lambda r, n: str(raw(r, n) if raw(r, n) is not None else '').strip()
        before = dict(self.rows); add = 0; upd = 0; today = date.today().isoformat(); QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            for r in rows[h + 1:]:
                no = g(r, 'رقم الملف')
                if not no: continue
                d = {'file_no': no, 'sender': g(r, 'الجهة المرسلة'), 'send_no': g(r, 'رقم الإرسال'), 'send_date': to_iso(raw(r, 'تاريخ الإرسال'))}
                if 'ملاحظات' in head: d['notes'] = g(r, 'ملاحظات')
                for k, n in enumerate(PROV):
                    if n in head:
                        on = truthy(raw(r, n)); d[f'p{k}'] = '1' if on else ''; d[f't{k}'] = today if on else ''
                ex = self.db.by_no(no); self.db.save(d, ex, commit=False)
                if ex: upd += 1
                else: add += 1
            self.db.c.commit()
        except sqlite3.Error as e:
            self.db.c.rollback(); QApplication.restoreOverrideCursor(); return QMessageBox.warning(self, 'خطأ', f'تعذّر الاستيراد:\n{e}')
        QApplication.restoreOverrideCursor(); self.reload(); self.record(f'استيراد إكسل ({add + upd} ملف)', [(i, before.get(i), self.rows.get(i)) for i in set(before) | set(self.rows)]); self.refresh(); QMessageBox.information(self, 'تم الاستيراد', f'ملفات جديدة: {add}\nملفات محدَّثة: {upd}\n(يمكنك التراجع بالسهم ↶)')

    def export_pdf(self):
        if not self.shown:
            return QMessageBox.information(self, 'تنبيه', 'لا توجد ملفات للتصدير.\nاستعمل البحث أو الفلاتر أولاً، مثلاً: ✘ لم تُجب: الجلفة.')
        d = QDialog(self, Qt.Window); d.setWindowTitle('🖨️ مراجعة وتعديل تقرير PDF قبل الطباعة'); d.resize(1500, 760); lay = QVBoxLayout(d)
        info = QLabel('<b>مراجعة التقرير قبل التصدير</b><br>يمكنك حذف أي صف من التقرير، وتعديل رقم الملف أو الجهة المرسلة أو رقم الإرسال. وعند اختيار محافظة، سيظهر في التقرير عمود تلك المحافظة فقط، وسيتم تصدير الصفوف المطابقة للتصفية الحالية.'); info.setWordWrap(True); lay.addWidget(info)
        tools = QHBoxLayout(); prov = QComboBox(); prov.addItems(['كل المحافظات', *[f'✘ لم تُجب: {x}' for x in PROV], *[f'✔ أجابت: {x}' for x in PROV]])
        if hasattr(self, 'fProv'): prov.setCurrentIndex(self.fProv.currentIndex())
        search = QLineEdit(); search.setPlaceholderText('🔍 بحث داخل التقرير: رقم الملف، الجهة، رقم الإرسال...'); count = QLabel(); count.setObjectName('hint'); tools.addWidget(QLabel('تصفية التقرير:')); tools.addWidget(prov, 1); tools.addWidget(search, 2); tools.addWidget(count); lay.addLayout(tools)
        def selected_prov_index():
            idx = prov.currentIndex();
            if 1 <= idx <= 5: return idx - 1
            if 6 <= idx <= 10: return idx - 6
            return None
        def report_headers():
            k = selected_prov_index();
            return (['العدد', 'رقم الملف', 'الجهة المرسلة', 'رقم الإرسال', PROV[k], 'الحالة'] if k is not None else PDF_HEAD)
        cols = report_headers(); t = QTableWidget(0, len(cols)); t.setHorizontalHeaderLabels(cols); t.setAlternatingRowColors(True); t.setSelectionBehavior(QAbstractItemView.SelectRows); t.setSelectionMode(QAbstractItemView.SingleSelection); t.setWordWrap(False); t.setLayoutDirection(Qt.RightToLeft); editable_cols = {1,2,3}; hh = t.horizontalHeader(); hh.setSectionResizeMode(QHeaderView.Interactive); hh.setStretchLastSection(True); hh.setMinimumHeight(44)
        def set_widths():
            widths = ([65, 125, 220, 115, 120, 145] if selected_prov_index() is not None else [65, 125, 220, 115, 100, 100, 100, 100, 100, 145])
            for j, w in enumerate(widths): t.setColumnWidth(j, w)
        set_widths(); lay.addWidget(t, 1)
        actions = QHBoxLayout(); delete_btn = QPushButton('🗑 حذف ال��ف من التقرير'); reset_btn = QPushButton('↺ إعادة التقرير الأصلي'); actions.addWidget(delete_btn); actions.addWidget(reset_btn); actions.addStretch(); cancel = QPushButton('إلغاء'); export = QPushButton('📄 إنشاء PDF'); export.setDefault(True); actions.addWidget(cancel); actions.addWidget(export); lay.addLayout(actions)
        original = [dict(r) for r in self.shown]; current = [dict(r) for r in self.shown]
        def selected_prov_match(r):
            idx = prov.currentIndex()
            if idx == 0: return True
            if 1 <= idx <= 5: return not bool(r.get(f'p{idx - 1}'))
            k = idx - 6; return bool(r.get(f'p{k}'))
        def matches(r):
            q = norm(search.text()).strip();
            if not selected_prov_match(r): return False
            if not q: return True
            haystack = norm(' '.join([str(r.get('file_no', '')), str(r.get('sender', '')), str(r.get('send_no', ''))]))
            return all(x in haystack for x in q.split())
        def fill():
            L = [r for r in current if matches(r)]; t.blockSignals(True); t.setUpdatesEnabled(False); t.clearContents(); t.setRowCount(len(L)); k_selected = selected_prov_index(); t.setColumnCount(len(report_headers())); t.setHorizontalHeaderLabels(report_headers()); set_widths();
            for i, r in enumerate(L):
                vals = [str(i+1), r.get('file_no', ''), r.get('sender', ''), r.get('send_no', '')]
                for j, v in enumerate(vals):
                    it = QTableWidgetItem(str(v)); it.setTextAlignment(Qt.AlignVCenter | Qt.AlignRight); it.setFlags(it.flags() | Qt.ItemIsEditable if j in editable_cols else it.flags() & ~Qt.ItemIsEditable); t.setItem(i, j, it)
                if k_selected is None: prov_indexes = list(range(5))
                else: prov_indexes = [k_selected]
                for pos, k in enumerate(prov_indexes):
                    ok = bool(r.get(f'p{k}')); it = QTableWidgetItem(YES if ok else NO); it.setTextAlignment(Qt.AlignCenter); it.setBackground(GREEN_BG if ok else RED_BG); it.setForeground(GREEN_FG if ok else RED_FG); it.setFont(self.big); it.setFlags(it.flags() & ~Qt.ItemIsEditable); t.setItem(i, 4 + pos, it)
                n = done(r); it = QTableWidgetItem('مكمل الإجابة ✔' if n == 5 else f'معلق لم يكتمل {n}/5'); it.setTextAlignment(Qt.AlignCenter); it.setBackground(DONE_BG if n == 5 else PEND_BG); it.setForeground(DONE_FG if n == 5 else PEND_FG); it.setFlags(it.flags() & ~Qt.ItemIsEditable); t.setItem(i, 5 + (0 if k_selected is not None else 5), it); t.item(i, 0).setData(Qt.UserRole, r.get('id'))
            t.setUpdatesEnabled(True); t.blockSignals(False); count.setText(f'📄 {len(L)} ملف في التقرير')
        def sync_edit(row, col):
            rid = t.item(row, 0).data(Qt.UserRole); r = next((x for x in current if x.get('id') == rid), None)
            if not r: return
            if col == 1: r['file_no'] = t.item(row, col).text().strip()
            elif col == 2: r['sender'] = t.item(row, col).text().strip()
            elif col == 3: r['send_no'] = t.item(row, col).text().strip()
        def delete_row():
            row = t.currentRow();
            if row < 0: return QMessageBox.information(d, 'تنبيه', 'حدد صفاً أولاً.')
            rid = t.item(row, 0).data(Qt.UserRole); r = next((x for x in current if x.get('id') == rid), None)
            if not r: return
            current.remove(r); fill()
        def reset_report():
            current[:] = [dict(r) for r in original]; fill()
        def make_pdf():
            export_rows = [r for r in current if matches(r)]
            if not export_rows: return QMessageBox.information(d, 'تنبيه', 'لا توجد ملفات مطابقة للتصفية الحالية.')
            selected_k = selected_prov_index(); selected_label = (PROV[selected_k] if selected_k is not None else 'كل المحافظات')
            p, _ = QFileDialog.getSaveFileName(d, 'حفظ تقرير PDF', f'تقرير_التحقيق_{date.today()}.pdf', 'PDF (*.pdf)')
            if not p: return
            QApplication.setOverrideCursor(Qt.WaitCursor)
            try:
                pr = a4_printer(False); pr.setOutputFormat(QPrinter.PdfFormat); pr.setOutputFileName(p); table_doc(export_rows, f'أرشيف التحقيق العقاري — ولاية الجلفة — {selected_label}', selected_k).print_(pr)
            except Exception as e:
                QApplication.restoreOverrideCursor(); return QMessageBox.warning(d, 'خطأ', f'تعذّر إنشاء PDF:\n{e}')
            finally:
                QApplication.restoreOverrideCursor()
            d.accept()
            try: os.startfile(p)
            except Exception: pass
        t.cellChanged.connect(sync_edit); prov.currentIndexChanged.connect(lambda _: fill()); search.textChanged.connect(lambda _: fill()); delete_btn.clicked.connect(delete_row); reset_btn.clicked.connect(reset_report); cancel.clicked.connect(d.reject); export.clicked.connect(make_pdf); fill(); d.exec_()

    def apply_theme(self, animate=True):
        app = QApplication.instance();
        if self.dark_mode:
            app.setStyleSheet(DARK_QSS)
            self.themeBtn.setText('☀️ الوضع النهاري')
        else:
            app.setStyleSheet(QSS)
            self.themeBtn.setText('🌙 الوضع الليلي')
        self.settings.setValue('dark_mode', self.dark_mode)
        self.bulkBtn.setStyleSheet('font-weight:bold;color:#d9eee2' if self.dark_mode else 'font-weight:bold;color:#123b2e')
        if animate:
            self.flash_message('تم التبديل إلى ' + ('الوضع الليلي 🌙' if self.dark_mode else 'الوضع النهاري ☀️'))

    def toggle_theme(self):
        self.dark_mode = not self.dark_mode; self.apply_theme(True)

    def flash_message(self, msg):
        if not hasattr(self, 'msg') or not self.msg: return
        self.msg.setText(msg); self.msg.show();
        if hasattr(self, '_msg_timer'): self._msg_timer.stop();
        self._msg_timer = QTimer(self); self._msg_timer.setSingleShot(True); self._msg_timer.timeout.connect(self.msg.hide); self._msg_timer.start(2200)

    def start_ui_animation(self):
        try:
            from PyQt5.QtWidgets import QGraphicsOpacityEffect
            effect = QGraphicsOpacityEffect(self); self.setGraphicsEffect(effect); anim = QPropertyAnimation(effect, b'opacity', self); anim.setDuration(520); anim.setStartValue(0.0); anim.setEndValue(1.0); anim.setEasingCurve(QEasingCurve.OutCubic); self._fade_anim = anim; anim.finished.connect(lambda: self.setGraphicsEffect(None)); anim.start()
        except Exception:
            pass

    def backup(self):
        d = QFileDialog.getExistingDirectory(self, 'اختر مكان حفظ النسخة الاحتياطية')
        if d:
            z = shutil.make_archive(os.path.join(d, f'نسخة_ارشيف_{date.today()}'), 'zip', DATA_DIR)
            QMessageBox.information(self, 'تم', f'تم حفظ النسخة الاحتياطية ✓\n{z}')


if __name__ == '__main__':
    app = QApplication(sys.argv)
    app.setLayoutDirection(Qt.RightToLeft)
    app.setFont(APP_FONT)
    app.setStyle('Fusion')
    app.setStyleSheet(QSS)
    app.setQuitOnLastWindowClosed(False)
    if LoginDlg(first=not has_pw()).exec_() != QDialog.Accepted:
        sys.exit(0)
    w = Main(); app.setQuitOnLastWindowClosed(True); w.showMaximized(); sys.exit(app.exec_())
